# -*- coding: utf-8 -*-
"""Тезисы для организатора — в Word.

Тот же текст, что в tezisy.md, но в формате, который открывается на
телефоне и пересылается в мессенджере не глядя. Разметку для
разработчиков организатору слать незачем: её просто не открыть.

Запуск:  python3 tezisy_docx.py
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

TUT = os.path.dirname(os.path.abspath(__file__))
VYHOD = os.path.join(TUT, '..', 'tezisy-dlya-pressy.docx')

SINIJ = RGBColor(0x1B, 0x35, 0x7A)
SERYJ = RGBColor(0x44, 0x4A, 0x5A)

EKSPERT = 'Руслан Анатольевич Дёмин'
DOLZHNOST = 'управляющий партнёр'

CITATY = [
    'Контент-фабрики есть у всех, это не новость. Новость в том, что '
    'управление циклом можно отдать самому ИИ: он лучше человека знает, '
    'по каким правилам его собратья отбирают контент.',

    'Предприниматель тратит на соцсети не деньги, а своё время — самый '
    'дорогой ресурс в компании. Цикл забирает рутину, человеку остаётся '
    'решение и смысл.',

    'Чтобы работать с нейросетями самому, нужно покупать токены, держать '
    'зарубежные аккаунты и воевать с блокировками. Готовый сервис снимает '
    'это целиком.',

    'Клиент больше не листает десять ссылок — он спрашивает нейросеть и '
    'получает два-три названия. Попасть в этот ответ можно '
    'целенаправленно: это такая же работа, как SEO, только по другим '
    'правилам.',

    'Попасть в нейроответ можно целенаправленно, но это не кнопка: от трёх '
    'до шести месяцев работы. Зато эту работу машина делает лучше '
    'человека — и без агентского счёта каждый месяц.',
]


def abzac(doc, tekst, razmer=11, zhirno=False, kursiv=False, cvet=None,
          otstup=0, pered=0, posle=6, vyravnivanie=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(pered)
    p.paragraph_format.space_after = Pt(posle)
    if otstup:
        p.paragraph_format.left_indent = Inches(otstup)
    if vyravnivanie is not None:
        p.alignment = vyravnivanie
    r = p.add_run(tekst)
    r.font.size = Pt(razmer)
    r.font.name = 'Calibri'
    r.bold = zhirno
    r.italic = kursiv
    if cvet is not None:
        r.font.color.rgb = cvet
    return p


def zagolovok(doc, tekst, razmer=14, pered=16):
    abzac(doc, tekst, razmer=razmer, zhirno=True, cvet=SINIJ,
          pered=pered, posle=4)


def main():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)

    abzac(doc, 'Тема и тезисы для пресс-конференции', razmer=18, zhirno=True,
          cvet=SINIJ, posle=2)
    abzac(doc, '«Эффективные инструменты привлечения клиентов» · '
               '20 октября 2026 · Москва, Кузнецкий мост, 7',
          razmer=10, cvet=SERYJ, posle=2)
    abzac(doc, 'Эксперт: %s, %s' % (EKSPERT, DOLZHNOST), razmer=10,
          cvet=SERYJ, posle=10)

    zagolovok(doc, 'Тема', pered=8)
    abzac(doc, 'Полный ИИ-цикл для бизнеса: от исследования до публикации. '
               'Исследование, стратегия, генерация текстов, картинок и видео, '
               'автопубликация на 6+ площадках. Продвижение в SEO и GEO.')

    zagolovok(doc, 'Три направления')
    for nomer, (imya, text) in enumerate((
        ('Как не тратить время на SMM.',
         ' Контент-фабрики есть у всех, но управляет ими человек: проверяет, '
         'правит, раскладывает по площадкам. Управление циклом можно отдать '
         'самому ИИ.'),
        ('Как сэкономить от 120–150 000 ₽ в месяц.',
         ' Без агентства, без покупки токенов, без зарубежных аккаунтов и '
         'борьбы с блокировками — готовый сервис за фиксированную плату при '
         'сравнимом эффекте.'),
        ('Как попасть в нейровыдачу.',
         ' Это не кнопка: от трёх до шести месяцев работы, и эту работу '
         'машина делает лучше человека.'),
    ), 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run('%d. %s' % (nomer, imya))
        r.bold = True
        r.font.size = Pt(11)
        r.font.name = 'Calibri'
        r = p.add_run(text)
        r.font.size = Pt(11)
        r.font.name = 'Calibri'

    zagolovok(doc, 'Готовые цитаты')
    abzac(doc, 'Можно ставить в текст как есть, с подписью «%s, %s».'
               % (EKSPERT.split()[0] + ' ' + EKSPERT.split()[2], DOLZHNOST),
          razmer=10, cvet=SERYJ, posle=8)
    for c in CITATY:
        abzac(doc, '«%s»' % c, razmer=11.5, kursiv=True, otstup=0.3, posle=10)

    zagolovok(doc, 'Справка о компаниях')
    abzac(doc, 'Для блока в конце публикации.', razmer=10, cvet=SERYJ,
          posle=8)
    for imya, text in (
        ('Промптер', 'российский сервис автоматизации SMM: исследование ниши, '
                     'стратегия, генерация текстов, изображений и видео, '
                     'публикация на 6+ площадках.'),
        ('ГенИИ', 'официальный партнёр «Промптера» по Москве и области. '
                  'Продвижение компаний в ответах нейросетей: SEO и GEO.'),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run('%s — ' % imya)
        r.bold = True
        r.font.size = Pt(11)
        r.font.name = 'Calibri'
        r = p.add_run(text)
        r.font.size = Pt(11)
        r.font.name = 'Calibri'

    zagolovok(doc, 'Контакты')
    abzac(doc, EKSPERT, razmer=11, zhirno=True, posle=2)
    abzac(doc, DOLZHNOST, razmer=10, cvet=SERYJ, posle=4)
    abzac(doc, '+7 906 758-77-77   ·   prompter.moscow@mail.ru\n'
          'prompter-ai.moscow   ·   genii-ai.ru',
          razmer=11)

    put = os.path.abspath(VYHOD)
    doc.save(put)
    print('собрано:', put)
    return put


if __name__ == '__main__':
    main()
