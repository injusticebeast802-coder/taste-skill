# -*- coding: utf-8 -*-
"""Инструменты сборки презентации: палитра сайта, шрифты, измерение текста.

Главная мысль файла: ширину каждой строки мы считаем сами, по тем же
файлам шрифтов, которыми набран сайт. Поэтому размер кегля не
подбирается на глаз — он выбирается как самый крупный из тех, при
которых текст гарантированно влезает в отведённую коробку.
"""
import os, tempfile
from PIL import Image, ImageDraw, ImageFont
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Пути считаем от этого файла: папка sborka лежит внутри репозитория,
# поэтому скрипт работает на любой машине без правок.
TUT   = os.path.dirname(os.path.abspath(__file__))
KOREN = os.path.abspath(os.path.join(TUT, '..', '..', '..'))
SAJT  = os.path.join(KOREN, 'prompter-moscow')

SHRIFTY = os.path.join(KOREN, 'analiz', 'shrifty')
IMG     = os.path.join(SAJT, 'img')
ASSETS  = os.path.join(SAJT, 'assets')
# Обрезанные картинки и значки — во временную папку, а не в репозиторий.
TMP     = os.path.join(tempfile.gettempdir(), 'prompter-preza', 'media')
os.makedirs(TMP, exist_ok=True)

# Палитра взята из style.css сайта, один в один.
BG     = RGBColor(0x0A, 0x0A, 0x0E)   # --bg
CARD   = RGBColor(0x15, 0x23, 0x43)   # --surface, синяя ячейка
CARD2  = RGBColor(0x1B, 0x2D, 0x55)   # --surface-2
LINE   = RGBColor(0x2C, 0x44, 0x78)   # --line поверх тёмного фона
LINE2  = RGBColor(0x3E, 0x5C, 0x9C)   # --line-2
BLUE   = RGBColor(0x3D, 0x6B, 0xFF)   # --accent
LBLUE  = RGBColor(0x7F, 0xC4, 0xFF)   # --accent-text, цифры и выделения
GREEN  = RGBColor(0x4A, 0xDE, 0x80)   # --green, выгода и экономия
RED    = RGBColor(0xE8, 0x45, 0x3C)   # --accent-2, только знак и редкие метки
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
MUTED  = RGBColor(0xD6, 0xE4, 0xFF)   # --muted
MUTED2 = RGBColor(0x9F, 0xB6, 0xDC)   # --muted-2

DISP = 'Unbounded'   # заголовки, как на сайте
BODY = 'Inter'       # весь остальной текст

W, H = 13.333, 7.5
PAD  = 0.62                 # поле страницы
CW   = W - PAD * 2          # рабочая ширина

_FILES = {
    ('Unbounded', True):  'Unbounded-Bold.ttf',
    ('Unbounded', False): 'Unbounded-SemiBold.ttf',
    ('Inter', 'r'):       'Inter-Regular.ttf',
    ('Inter', 'm'):       'Inter-Medium.ttf',
    ('Inter', 's'):       'Inter-Bold.ttf',
    ('Inter', 'b'):       'Inter-Bold.ttf',
}
_cache = {}


def _font(key, pt):
    """PIL-шрифт увеличенного кегля: так измерение точнее."""
    k = (key, round(pt, 2))
    if k not in _cache:
        _cache[k] = ImageFont.truetype(os.path.join(SHRIFTY, _FILES[key]), int(round(pt * 4)))
    return _cache[k]


def shirina(text, key, pt):
    """Ширина строки в дюймах при заданном кегле."""
    return _font(key, pt).getlength(text) / 4.0 / 72.0


def perenos(text, key, pt, maxw):
    """Разбивка на строки по ширине коробки. Возвращает список строк."""
    out = []
    for kus in text.split('\n'):
        slova, stroka = kus.split(' '), ''
        for sl in slova:
            proba = sl if not stroka else stroka + ' ' + sl
            if shirina(proba, key, pt) <= maxw or not stroka:
                stroka = proba
            else:
                out.append(stroka)
                stroka = sl
        out.append(stroka)
    return out


def podobrat(text, key, kegli, maxw, maxh, lead, max_strok=None):
    """Самый крупный кегль, при котором текст влезает в коробку.

    Проверяем три вещи: ни одна строка не шире коробки (иначе длинное
    слово вроде «Автопубликация» разорвётся посередине), строк не больше
    разрешённого и высота не больше отведённой.

    Запас 4 процента по ширине: PowerPoint переносит строки по своим
    правилам, и без запаса пограничная строка у клиента уедет вниз.
    """
    for pt in kegli:
        lines = perenos(text, key, pt, maxw * 0.96)
        if any(shirina(l, key, pt) > maxw for l in lines):
            continue                      # слово не помещается по ширине
        if max_strok and len(lines) > max_strok:
            continue
        if len(lines) * pt * lead / 72.0 <= maxh:
            return pt, lines
    pt = kegli[-1]
    return pt, perenos(text, key, pt, maxw * 0.96)


# ---------------------------------------------------------------- картинки

def kadr(src, w_in, h_in, dst, dark=0.0, radius=0.0, left_wash=0.0):
    """Кадрирует картинку под коробку без искажений, скругляет углы.

    dark       — общее затемнение, 0..1
    left_wash  — затемнение слева направо: под текст на обложке
    radius     — радиус скругления в дюймах
    """
    im = Image.open(src).convert('RGB')
    tw, th = int(w_in * 200), int(h_in * 200)
    k = max(tw / im.width, th / im.height)
    im = im.resize((max(1, int(im.width * k)), max(1, int(im.height * k))), Image.LANCZOS)
    x, y = (im.width - tw) // 2, (im.height - th) // 2
    im = im.crop((x, y, x + tw, y + th))

    if dark > 0:
        im = Image.blend(im, Image.new('RGB', im.size, (10, 10, 14)), dark)
    if left_wash > 0:
        maska = Image.new('L', im.size)
        d = ImageDraw.Draw(maska)
        for i in range(im.width):
            d.line([(i, 0), (i, im.height)],
                   fill=int(255 * left_wash * max(0.0, 1.0 - (i / im.width) / 0.72)))
        im = Image.composite(Image.new('RGB', im.size, (10, 10, 14)), im, maska)

    im = im.convert('RGBA')
    if radius > 0:
        r = int(radius * 200)
        a = Image.new('L', im.size, 0)
        ImageDraw.Draw(a).rounded_rectangle([0, 0, im.width - 1, im.height - 1], r, fill=255)
        im.putalpha(a)
    put = os.path.join(TMP, dst)
    im.save(put)
    return put


# ---------------------------------------------------------------- фигуры

def fon(sl):
    """Сплошной тёмный фон слайда."""
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(H))
    s.fill.solid(); s.fill.fore_color.rgb = BG
    s.line.fill.background(); s.shadow.inherit = False
    return s


def plitka(sl, x, y, w, h, fill=CARD, line=LINE, radius=0.16, lw=1.0):
    """Карточка сайта: скруглённый прямоугольник с тонкой синей рамкой."""
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.adjustments[0] = min(0.5, radius / min(w, h))
    s.shadow.inherit = False
    return s


def nadpis(sl, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def abzac(tf, kuski, size, font=BODY, color=WHITE, bold=False, lead=1.4,
          space_before=0, align=PP_ALIGN.LEFT, first=False):
    """Абзац из кусков. Кусок — строка или (текст, {цвет, жирность}).

    Куски нужны, чтобы внутри одной фразы выделить слово цветом, как на
    сайте: «Не просто генерирует — сама публикует».
    """
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = lead
    if space_before:
        p.space_before = Pt(space_before)
    for kus in (kuski if isinstance(kuski, list) else [kuski]):
        txt, opt = (kus, {}) if isinstance(kus, str) else kus
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size)
        r.font.name = opt.get('font', font)
        r.font.bold = opt.get('bold', bold)
        r.font.color.rgb = opt.get('color', color)
    return p


def kartinka(sl, put, x, y, w, h):
    return sl.shapes.add_picture(put, Inches(x), Inches(y), Inches(w), Inches(h))


def ssylka(shape, url):
    shape.click_action.hyperlink.address = url
