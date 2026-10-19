# -*- coding: utf-8 -*-
"""Сборка конференционной версии прямо в файле prompter-smm-2026.pptx.

Раньше листы отдавались отдельным файлом и вставлялись руками. Теперь
сама презентация лежит в репозитории, и листы можно поставить на свои
места: лишнего копирования нет, порядок задан здесь и не разъедется.

Что делаем:
  * «Темы для конференции» — вторым листом и повтором в конец;
  * четыре листа доводов — после тарифов, перед кейсами;
  * лист контактов меняем на «Полный AI цикл» с кодом и контактами
    с визитки: строка «Ваш менеджер» в конференционной версии лишняя.

Две тонкости, из-за которых это не просто «добавить слайды».

Размер листа у презентации 10 × 5,625 дюйма, а вся наша вёрстка
считалась под 13,333 × 7,5. Пересчитывать каждое число руками — верный
способ что-нибудь пропустить, поэтому подменяются сами Inches и Pt:
вёрстка остаётся в прежней сетке, а на лист ложится уменьшенной.

Фон и знак «Промптера» берём из самой презентации, со слайда
контактов. Нарисованный заново фон был похож, но здесь он настоящий.
"""
import os
import sys

TUT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(TUT), 'sborka'))
sys.path.insert(0, TUT)

from pptx import Presentation
from pptx.util import Inches as _Inches, Pt as _Pt

PREZA = os.path.join(os.path.dirname(TUT), 'prompter-smm-2026.pptx')
VYHOD = os.path.join(os.path.dirname(TUT),
                     'prompter-smm-2026-konferenciya.pptx')

# Отношение сторон у обоих размеров одно и то же, поэтому хватает
# одного множителя: 10 / 13.333.
MASHTAB = 0.75


# Шрифт намеренно Arial, а не наши Unbounded и Inter. В самой
# презентации весь живой текст набран им же, шрифты внутрь файла
# не вложены, а показывать будут с чужого компьютера. Шрифт, которого
# там нет, PowerPoint подменит своим — и поедут и кегли, и переносы.
SHRIFT = 'Arial'

# Ширину строки считаем по Liberation Sans: он метрически совпадает
# с Arial, то есть кеголь подбирается по тому, чем текст в итоге и наберётся.
LIB = '/usr/share/fonts/truetype/liberation'
PRYAMOJ = os.path.join(LIB, 'LiberationSans-Regular.ttf')
ZHIRNYJ = os.path.join(LIB, 'LiberationSans-Bold.ttf')


def nastroit_masshtab():
    """Подменяет единицы измерения и шрифты в вёрстке."""
    import kit
    import oranzhevye

    def dyujmy(v):
        return _Inches(v * MASHTAB)

    def punkty(v):
        return _Pt(v * MASHTAB)

    for modul in (kit, oranzhevye):
        modul.Inches = dyujmy
        modul.Pt = punkty
        modul.DISP = SHRIFT
        modul.BODY = SHRIFT

    # Заголовки и полужирные строки меряем по жирному начертанию:
    # в одной строке часто соседствуют оба, и лучше перестраховаться.
    kit._FILES = {
        ('Unbounded', True): ZHIRNYJ,
        ('Unbounded', False): ZHIRNYJ,
        ('Inter', 'r'): PRYAMOJ,
        ('Inter', 'm'): ZHIRNYJ,
        ('Inter', 's'): ZHIRNYJ,
        ('Inter', 'b'): ZHIRNYJ,
    }
    kit.SHRIFTY = ''
    kit._cache.clear()
    return oranzhevye


def vytashchit_kartinki(prs):
    """Фон и знак «Промптера» — из слайда контактов самой презентации."""
    kontakty = prs.slides[13]
    kartinki = [sh for sh in kontakty.shapes if sh.shape_type == 13]
    kartinki.sort(key=lambda sh: sh.width, reverse=True)
    puti = []
    for sh, imya in zip(kartinki, ('fon-nastoyashchij.png', 'znak-prompter.png')):
        put = os.path.join(TUT, imya)
        with open(put, 'wb') as f:
            f.write(sh.image.blob)
        puti.append(put)
    return puti


def poryadok(prs, listy):
    """Ставит листы в заданный порядок и выкидывает остальные.

    Раньше здесь были перестановки по номерам — и после первого же
    удалённого листа все номера ниже съехали на единицу, и листы
    разошлись по чужим местам. Поэтому порядок задаётся сразу
    целиком, самими листами, а не шагами по номерам.
    """
    spisok = prs.slides._sldIdLst
    for el in list(spisok):
        spisok.remove(el)
    for el in listy:
        spisok.append(el)


def main():
    prs = Presentation(PREZA)
    bylo = list(prs.slides._sldIdLst)
    print('в презентации листов:', len(bylo),
          '| размер %.2f × %.2f' % (prs.slide_width / 914400,
                                    prs.slide_height / 914400))

    fon, znak = vytashchit_kartinki(prs)
    print('фон и знак вынуты из слайда контактов')

    o = nastroit_masshtab()
    o.FON = fon
    o.PROMPTER_ZNAK = znak
    o.STRANICA = 0
    o.VSEGO = 0          # номера страниц рисует сама презентация

    def dobavit(risovat):
        risovat(prs)
        return list(prs.slides._sldIdLst)[-1]

    temy_1 = dobavit(o.slajd_temy_konf)
    fabriki = dobavit(o.slajd_3)
    vremya = dobavit(o.slajd_4)
    dengi = dobavit(o.slajd_5)
    vydacha = dobavit(o.slajd_6)
    cikl = dobavit(o.slajd_8)
    temy_2 = dobavit(o.slajd_temy_konf)

    # Собственный лист контактов (14-й) в этот порядок не входит:
    # на нём стояла строка «Ваш менеджер», а это презентация
    # конференции, а не конкретного менеджера.
    poryadok(prs, [
        bylo[0],                       # 1  обложка со слоганом
        temy_1,                        # 2  темы для конференции
    ] + bylo[1:11] + [                 # 3–12 платформа и тарифы
        fabriki,                       # 13 контент-фабрики
        vremya,                        # 14 время
        dengi,                         # 15 деньги
        vydacha,                       # 16 нейровыдача
    ] + bylo[11:13] + [                # 17–18 кейсы
        temy_2,                        # 19 темы ещё раз
        cikl,                          # 20 полный AI цикл, код и контакты
    ])
    print('листов стало:', len(prs.slides))

    if o.MALO:
        print('НЕ ВЛЕЗЛО:')
        for m in o.MALO:
            print('  -', m)
    else:
        print('всё влезает в отведённое место')

    prs.save(VYHOD)
    print('собрано:', VYHOD)
    return VYHOD


if __name__ == '__main__':
    main()
